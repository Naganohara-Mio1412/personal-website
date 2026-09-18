"""Generate cleaned data workbooks, CSV exports, and formatted Excel reports."""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from profile_data import build_report_tables

LOGGER = logging.getLogger(__name__)
INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
INVALID_FILE_CHARS = re.compile(r"[<>:\"/\\|?*]")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _unique_excel_names(names: Iterable[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    used: set[str] = set()
    for original in names:
        base = INVALID_SHEET_CHARS.sub("_", str(original)).strip(" '") or "Dataset"
        base = base[:31]
        candidate = base
        counter = 2
        while candidate.casefold() in used:
            suffix = f"_{counter}"
            candidate = f"{base[:31 - len(suffix)]}{suffix}"
            counter += 1
        used.add(candidate.casefold())
        mapping[original] = candidate
    return mapping


def _csv_component(name: str) -> str:
    value = INVALID_FILE_CHARS.sub("_", str(name)).strip(" .") or "dataset"
    return value[:80]


def _format_worksheet(sheet: Worksheet) -> None:
    """Apply common readable formatting to a populated worksheet."""
    sheet.freeze_panes = "A2"
    if sheet.max_row >= 1 and sheet.max_column >= 1:
        sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    headers = {cell.column: str(cell.value or "").casefold() for cell in sheet[1]}
    for column_cells in sheet.columns:
        column_index = column_cells[0].column
        header = headers.get(column_index, "")
        max_length = min(
            50,
            max(10, max((len(str(cell.value)) if cell.value is not None else 0) for cell in column_cells) + 2),
        )
        sheet.column_dimensions[column_cells[0].column_letter].width = max_length
        for cell in column_cells[1:]:
            if cell.value is None:
                continue
            if "ratio" in header or "percentage" in header or "percent" in header:
                cell.number_format = "0.00%"
            elif cell.is_date or "date" in header or "timestamp" in header:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
            elif isinstance(cell.value, int):
                cell.number_format = "#,##0"
            cell.alignment = Alignment(vertical="top")


def write_cleaned_data(datasets: dict[str, pd.DataFrame], output_dir: Path) -> list[Path]:
    """Write full cleaned data to one workbook and one or more UTF-8-SIG CSV files."""
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook_path = output_dir / "cleaned_data.xlsx"
    sheet_names = _unique_excel_names(datasets.keys())
    with pd.ExcelWriter(
        workbook_path,
        engine="openpyxl",
        datetime_format="yyyy-mm-dd hh:mm:ss",
        date_format="yyyy-mm-dd",
    ) as writer:
        for name, frame in datasets.items():
            frame.to_excel(writer, sheet_name=sheet_names[name], index=False)
        for sheet in writer.book.worksheets:
            _format_worksheet(sheet)

    created = [workbook_path]
    if len(datasets) == 1:
        csv_path = output_dir / "cleaned_data.csv"
        next(iter(datasets.values())).to_csv(
            csv_path, index=False, encoding="utf-8-sig", date_format="%Y-%m-%d %H:%M:%S"
        )
        created.append(csv_path)
    else:
        used: set[str] = set()
        for name, frame in datasets.items():
            base = _csv_component(name)
            candidate = base
            counter = 2
            while candidate.casefold() in used:
                candidate = f"{base}_{counter}"
                counter += 1
            used.add(candidate.casefold())
            csv_path = output_dir / f"cleaned_data__{candidate}.csv"
            frame.to_csv(
                csv_path, index=False, encoding="utf-8-sig", date_format="%Y-%m-%d %H:%M:%S"
            )
            created.append(csv_path)
    LOGGER.info("Generated cleaned workbook and %d CSV export(s)", len(created) - 1)
    return created


def write_cleaning_report(
    before: dict[str, pd.DataFrame],
    after: dict[str, pd.DataFrame],
    logs: list[dict[str, object]],
    output_path: Path,
) -> Path:
    """Write the eight-section formatted Excel cleaning report."""
    tables = build_report_tables(before, after, logs)
    with pd.ExcelWriter(
        output_path,
        engine="openpyxl",
        datetime_format="yyyy-mm-dd hh:mm:ss",
        date_format="yyyy-mm-dd",
    ) as writer:
        for sheet_name, table in tables.items():
            if table.empty and len(table.columns) == 0:
                table = pd.DataFrame({"status": ["No applicable findings"]})
            table.to_excel(writer, sheet_name=sheet_name, index=False)
        for sheet in writer.book.worksheets:
            _format_worksheet(sheet)
    LOGGER.info("Generated Excel cleaning report: %s", output_path)
    return output_path
