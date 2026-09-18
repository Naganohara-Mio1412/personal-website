"""End-to-end tests for safe cleaning, reports, and offline dashboard generation."""

from __future__ import annotations

import ast
import hashlib
import importlib
import json
import re
import subprocess
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from create_test_data import build_dirty_frame, create_test_data

SKILL_DIR = Path(__file__).resolve().parents[1]
PIPELINE = SKILL_DIR / "scripts" / "run_pipeline.py"
REQUIRED_REPORT_SHEETS = [
    "数据概览", "字段字典", "缺失值分析", "重复值分析", "异常值提示",
    "类别字段分析", "清洗前后对比", "清洗日志",
]
LOG_KEYS = {
    "timestamp", "dataset", "column", "action", "reason", "affected_rows",
    "before_summary", "after_summary", "automatic", "mode",
}

sys.path.insert(0, str(SKILL_DIR / "scripts"))
from load_data import load_data  # noqa: E402
from clean_data import clean_dataset  # noqa: E402

SCRIPT_MODULES = (
    "profile_data", "clean_data", "load_data", "generate_excel",
    "generate_dashboard", "run_pipeline",
)


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    digest.update(path.read_bytes())
    return digest.hexdigest()


def invoke_pipeline(source: Path, output: Path, cwd: Path) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(PIPELINE), "--input", str(source), "--output", str(output), "--mode", "safe"],
        cwd=cwd,
        text=True,
        encoding="utf-8",
        capture_output=True,
        check=False,
    )


@pytest.fixture()
def completed_run(tmp_path: Path) -> dict[str, object]:
    source = create_test_data(tmp_path / "input" / "dirty.csv")
    output = tmp_path / "result"
    before_hash = file_hash(source)
    before_mtime = source.stat().st_mtime_ns
    result = invoke_pipeline(source, output, tmp_path)
    assert result.returncode == 0, result.stderr
    return {
        "source": source,
        "output": output,
        "before_hash": before_hash,
        "before_mtime": before_mtime,
        "result": result,
    }


def test_original_file_is_unchanged(completed_run: dict[str, object]) -> None:
    source = completed_run["source"]
    assert isinstance(source, Path)
    assert file_hash(source) == completed_run["before_hash"]
    assert source.stat().st_mtime_ns == completed_run["before_mtime"]


def test_safe_cleaning_rules(completed_run: dict[str, object]) -> None:
    output = completed_run["output"]
    assert isinstance(output, Path)
    cleaned = pd.read_csv(
        output / "cleaned_data.csv",
        encoding="utf-8-sig",
        dtype={"Customer ID": "string"},
    )
    assert len(cleaned) == 5, "one empty row and one full duplicate must be removed"
    assert not cleaned.duplicated().any(), "fully duplicate rows must be removed"
    assert not cleaned.isna().all(axis=1).any(), "fully empty rows must be removed"
    assert "completely_empty" not in cleaned.columns, "fully empty columns must be removed"
    assert all(column == column.strip() for column in cleaned.columns)
    assert cleaned["Name"].tolist() == ["Alice", "Bob", "Carol", "Dana", "Eva"]
    assert cleaned["Customer ID"].tolist()[0] == "001", "identifier leading zeros must survive"
    assert cleaned["amount"].tolist()[0] == 1200.5, "unambiguous numbers must be converted"
    assert pd.to_datetime(cleaned["order_date"], errors="raise").notna().all()
    assert cleaned["active"].dtype == bool, "explicit boolean tokens must be normalized"
    assert cleaned["completion_percent"].tolist()[0] == "10%", "units must not be rescaled"
    assert 999999 in cleaned["score"].tolist(), "safe mode must retain statistical outliers"
    assert cleaned["business_note"].isna().sum() >= 2, "safe mode must not impute business missing values"
    assert "Beijing" in cleaned["category"].tolist()
    assert "beijing" in cleaned["category"].tolist(), "safe mode must not merge similar categories"


def test_ambiguous_binary_numbers_are_not_guessed_as_booleans() -> None:
    source = pd.DataFrame({"binary_code": ["0", "1", "0"]})
    cleaned, logs = clean_dataset("binary", source, mode="safe")
    assert str(cleaned["binary_code"].dtype) == "Int64"
    assert not any(entry["action"] == "convert_boolean" for entry in logs)


def test_logs_and_output_files(completed_run: dict[str, object]) -> None:
    output = completed_run["output"]
    assert isinstance(output, Path)
    required = [
        output / "cleaned_data.xlsx", output / "cleaned_data.csv",
        output / "cleaning_report.xlsx", output / "interactive_dashboard.html",
        output / "cleaning_log.json",
    ]
    assert all(path.is_file() and path.stat().st_size > 0 for path in required)
    logs = json.loads((output / "cleaning_log.json").read_text(encoding="utf-8"))
    assert logs
    assert all(LOG_KEYS.issubset(entry) for entry in logs)
    actions = {entry["action"] for entry in logs}
    assert {"remove_empty_rows", "remove_empty_column", "remove_duplicate_rows"}.issubset(actions)


def test_excel_report_and_outputs_are_readable(completed_run: dict[str, object]) -> None:
    output = completed_run["output"]
    assert isinstance(output, Path)
    cleaned_book = pd.ExcelFile(output / "cleaned_data.xlsx", engine="openpyxl")
    assert cleaned_book.sheet_names == ["dirty"]
    assert not pd.read_excel(output / "cleaned_data.xlsx", engine="openpyxl").empty
    assert not pd.read_csv(output / "cleaned_data.csv", encoding="utf-8-sig").empty
    report = load_workbook(output / "cleaning_report.xlsx", read_only=False)
    assert report.sheetnames == REQUIRED_REPORT_SHEETS
    for worksheet in report.worksheets:
        assert worksheet.freeze_panes == "A2"
        assert worksheet.auto_filter.ref


def test_dashboard_is_offline_and_interactive(completed_run: dict[str, object]) -> None:
    output = completed_run["output"]
    assert isinstance(output, Path)
    html = (output / "interactive_dashboard.html").read_text(encoding="utf-8")
    assert "Plotly.newPlot" in html
    assert "搜索字段..." in html
    assert "完整数据保存在 cleaned_data.xlsx 和 CSV 文件中" in html
    remote_resource = re.compile(
        r"<(?:script|link|img|iframe|source|video|audio|form)\b[^>]*"
        r"(?:src|href|action|poster)\s*=\s*['\"](?:https?:)?//",
        re.I,
    )
    assert not remote_resource.search(html)
    assert not re.search(r"(?:@import\s+|url\s*\(\s*)['\"]?(?:https?:)?//", html, re.I)
    assert not re.search(r"(?:fetch|importScripts)\s*\(\s*['\"]https?://", html, re.I)
    assert "cdn.plot.ly" not in html.casefold()
    assert not re.search(r"https?://", html, re.I)
    assert "default-src 'none'" in html
    assert "connect-src 'none'" in html, "the offline report must block network requests"
    assert not re.search(r"<script\b[^>]+src\s*=", html, re.I), "Plotly must be embedded, not linked"


def test_dashboard_modern_ui_contract(completed_run: dict[str, object]) -> None:
    output = completed_run["output"]
    assert isinstance(output, Path)
    html = (output / "interactive_dashboard.html").read_text(encoding="utf-8")

    # Compact product-style navigation and independently visible modules.
    assert 'role="tablist"' in html
    assert all(
        f'data-tab="{panel}"' in html
        for panel in ("overviewPanel", "fieldPanel", "relationPanel", "tablePanel")
    )
    assert 'class="field-layout"' in html
    assert 'id="fieldSearch"' in html
    assert 'id="numericChartTabs"' in html

    # Frontend-only date candidates and the three supported relationship modes.
    assert "function parseDashboardDate" in html
    assert "parsed/values.length>=.8" in html
    assert "NUMBER × NUMBER" in html
    assert "DATE × NUMBER" in html
    assert "CATEGORY × NUMBER" in html
    assert 'id="relationX"' in html and 'id="relationY"' in html

    # Filtering, debounced search, configurable pagination, and compact missingness UI.
    assert "debounce(applyFilters)" in html
    assert 'id="filterCondition"' in html
    assert 'id="filterValueMax"' in html
    assert 'id="pageSize"' in html
    assert 'class="progress-bar"' in html
    assert "type:'indicator'" not in html


def test_all_python_modules_import_and_no_absolute_paths() -> None:
    for module_name in SCRIPT_MODULES:
        assert importlib.import_module(module_name)
    python_files = list((SKILL_DIR / "scripts").glob("*.py")) + list(
        (SKILL_DIR / "tests").glob("*.py")
    )
    for path in python_files:
        tree = ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
        for node in ast.walk(tree):
            if isinstance(node, ast.Constant) and isinstance(node.value, str):
                assert not re.match(r"^[A-Za-z]:[\\/]", node.value), f"absolute path in {path}"
                assert not re.match(r"^/(?:home|Users|tmp|var)/", node.value), f"absolute path in {path}"


def test_multi_sheet_excel_is_not_merged_or_truncated(tmp_path: Path) -> None:
    source = tmp_path / "multi.xlsx"
    frame = create_test_data(tmp_path / "scratch.csv")
    first = pd.read_csv(frame, encoding="utf-8-sig")
    second = first.iloc[:3].copy()
    with pd.ExcelWriter(source, engine="openpyxl") as writer:
        first.to_excel(writer, sheet_name="North", index=False)
        second.to_excel(writer, sheet_name="South", index=False)
    output = tmp_path / "multi-output"
    result = invoke_pipeline(source, output, tmp_path)
    assert result.returncode == 0, result.stderr
    assert pd.ExcelFile(output / "cleaned_data.xlsx", engine="openpyxl").sheet_names == ["North", "South"]
    assert (output / "cleaned_data__North.csv").is_file()
    assert (output / "cleaned_data__South.csv").is_file()
    assert not (output / "cleaned_data.csv").exists()


def test_csv_tsv_and_json_loaders_and_chinese_encoding(tmp_path: Path) -> None:
    frame = build_dirty_frame().iloc[:3]
    gbk_csv = tmp_path / "中文数据.csv"
    tsv = tmp_path / "sample.tsv"
    json_path = tmp_path / "sample.json"
    frame.to_csv(gbk_csv, index=False, encoding="gbk")
    frame.to_csv(tsv, index=False, sep="\t", encoding="utf-8-sig")
    json_path.write_text(frame.to_json(orient="records", force_ascii=False), encoding="utf-8")
    assert next(iter(load_data(gbk_csv).values())).shape == frame.shape
    assert next(iter(load_data(tsv).values())).shape == frame.shape
    assert next(iter(load_data(json_path).values())).shape == frame.shape


def test_failure_returns_nonzero_without_success_artifacts(tmp_path: Path) -> None:
    missing = tmp_path / "does-not-exist.csv"
    output = tmp_path / "failed-output"
    result = invoke_pipeline(missing, output, tmp_path)
    assert result.returncode != 0
    assert "Pipeline failed" in result.stderr
    assert not output.exists() or not any(output.iterdir())


def test_output_name_conflict_never_overwrites_source(tmp_path: Path) -> None:
    source = create_test_data(tmp_path / "cleaned_data.csv")
    original_hash = file_hash(source)
    result = invoke_pipeline(source, tmp_path, tmp_path)
    assert result.returncode != 0
    assert file_hash(source) == original_hash
    assert not (tmp_path / "cleaning_report.xlsx").exists()
